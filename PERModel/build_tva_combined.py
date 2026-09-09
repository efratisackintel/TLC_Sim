"""Generate ONE combined TVA sim-vs-lab workbook covering all bandwidths.

Reuses the simulation engine, styling and chart logic from build_tva_compare and
writes a single Results/tva_sim_vs_lab_ALL.xlsx with, per bandwidth (20/40/80/160/320):
  - "TVA <bw>MHz"    : per-step table (sim/PR blue, genie purple, lab green)
  - "Charts <bw>MHz" : 5 line charts (throughput, MCS, BW, GI, streams)
  - "Data <bw>MHz"   : hidden chart-data source
plus a shared "RSSI to SNR" worked example. Regenerates from source data, so it does
not need (or touch) the per-bandwidth .xlsx files.

Usage:  python PERModel/build_tva_combined.py
"""
import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.marker import Marker

from build_tva_compare import (
    OUT, BANDWIDTHS, run_bw, nssTxt, giTxt, FILL, FCOL, BORD,
)

TVA_COLS = [
    ('#', 'd', lambda r: r['step']), ('SNR (dB)', 'd', lambda r: r['snr']),
    ('avg tpt (Mbps)', 's', lambda r: round(r['prTpt'])), ('PR MCS', 's', lambda r: r['prMcs']),
    ('PR BW', 's', lambda r: r['prBw']), ('PR streams', 's', lambda r: nssTxt(r['prNss'])),
    ('PR GI', 's', lambda r: giTxt(r['prGi08'])), ('PR SR', 's', lambda r: f"{round(r['prSR']*100)}%"),
    ('PR %time', 's', lambda r: f"{r['prPct']}%"),
    ('genie tpt (Mbps)', 'g', lambda r: round(r['geTpt'])), ('genie MCS', 'g', lambda r: r['geMcs']),
    ('genie BW', 'g', lambda r: r['geBw']), ('genie streams', 'g', lambda r: nssTxt(r['geNss'])),
    ('genie GI', 'g', lambda r: giTxt(r['geGi08'])),
    ('log RSSI', 'l', lambda r: round(r['lab']['rssi'], 1) if r['lab'] else ''),
    ('log tpt (Mbps)', 'l', lambda r: round(r['lab']['tpt']) if r['lab'] and r['lab']['tpt'] is not None else ''),
    ('log MCS', 'l', lambda r: r['lab']['mcs'] if r['lab'] and r['lab']['mcs'] is not None else ''),
    ('log BW', 'l', lambda r: r['lab']['bw'] if r['lab'] and r['lab']['bw'] is not None else ''),
    ('log streams', 'l', lambda r: nssTxt(r['lab']['nss']) if r['lab'] else ''),
    ('log GI', 'l', lambda r: giTxt(r['lab']['gi08']) if r['lab'] and r['lab']['gi08'] is not None else ''),
    ('log SR', 'l', lambda r: f"{round(r['lab']['sr']*100)}%" if r['lab'] and r['lab']['sr'] is not None else ''),
]
COL_WIDTHS = [5, 9, 14, 8, 7, 11, 8, 7, 9, 16, 10, 9, 13, 9, 9, 14, 9, 8, 11, 8, 8]


def write_tva_sheet(ws, bw, rows, off, lab_src):
    ws.append([f'{bw} MHz  |  Lab log source: {lab_src}    |    '
               f'RSSI\u2192SNR offset (MCS-knee aligned): {round(off,1)} dB'])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(TVA_COLS))
    ws.cell(1, 1).font = Font(italic=True, color='FF1F6B45')
    ws.append([c[0] for c in TVA_COLS])
    for i, c in enumerate(TVA_COLS, 1):
        cell = ws.cell(2, i)
        cell.fill = PatternFill('solid', fgColor=FILL.get(c[1] + 'h', FILL['dh']))
        cell.font = Font(bold=True, color='FF' + FCOL.get(c[1], FCOL['d']))
        cell.alignment = Alignment('center', 'center', wrap_text=True)
        cell.border = BORD
    for r in rows:
        ws.append([c[2](r) for c in TVA_COLS])
        rr = ws.max_row
        for i, c in enumerate(TVA_COLS, 1):
            cell = ws.cell(rr, i)
            cell.alignment = Alignment(horizontal='center')
            cell.border = BORD
            if c[1] != 'd':
                cell.fill = PatternFill('solid', fgColor=FILL[c[1]])
            cell.font = Font(color='FF' + FCOL.get(c[1], FCOL['d']))
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = 'A3'


def write_charts(wb, bw, rows):
    cd = wb.create_sheet(f'Data {bw}MHz')
    cd.sheet_state = 'hidden'
    cd.append(['SNR', 'sim tpt', 'genie tpt', 'lab tpt', 'sim MCS', 'genie MCS', 'lab MCS',
               'sim BW', 'genie BW', 'lab BW', 'sim GI', 'genie GI', 'lab GI',
               'sim streams', 'genie streams', 'lab streams'])
    for r in rows:
        L = r['lab']
        cd.append([r['snr'], round(r['prTpt']), round(r['geTpt']),
                   (round(L['tpt']) if L and L['tpt'] is not None else None),
                   r['prMcs'], r['geMcs'], (L['mcs'] if L else None),
                   r['prBw'], r['geBw'], (L['bw'] if L else None),
                   (0.8 if r['prGi08'] else 3.2), (0.8 if r['geGi08'] else 3.2),
                   ((0.8 if L['gi08'] else 3.2) if L and L['gi08'] is not None else None),
                   r['prNss'], r['geNss'], (L['nss'] if L else None)])
    n = len(rows)
    ch_ws = wb.create_sheet(f'Charts {bw}MHz')
    cats = Reference(cd, min_col=1, min_row=2, max_row=1 + n)
    syms = ['x', 'square', 'circle']
    lines = ['0071C5', '7C3AED', '2E8B57']

    def mkchart(title, c3, anchor):
        ch = LineChart(); ch.title = title; ch.height = 8.5; ch.width = 20
        ch.x_axis.title = 'SNR (dB)'; ch.x_axis.delete = False; ch.y_axis.delete = False
        for col in c3:
            ch.add_data(Reference(cd, min_col=col, min_row=1, max_row=1 + n), titles_from_data=True)
        ch.set_categories(cats)
        for si, ser in enumerate(ch.series):
            ser.graphicalProperties.line.solidFill = lines[si]
            ser.graphicalProperties.line.width = 20000
            ser.marker = Marker(symbol=syms[si], size=6); ser.smooth = False
        ch_ws.add_chart(ch, anchor)

    mkchart(f'{bw} MHz — Throughput (Mbps) — sim vs genie vs lab', [2, 3, 4], 'A1')
    mkchart(f'{bw} MHz — MCS — sim vs genie vs lab', [5, 6, 7], 'A18')
    mkchart(f'{bw} MHz — Bandwidth (MHz) — sim vs genie vs lab', [8, 9, 10], 'A35')
    mkchart(f'{bw} MHz — Guard interval (\u00b5s) — sim vs genie vs lab', [11, 12, 13], 'A52')
    mkchart(f'{bw} MHz — Streams (SISO=1 / MIMO=2) — sim vs genie vs lab', [14, 15, 16], 'A69')


def write_rssi_sheet(wb):
    rs = wb.create_sheet('RSSI to SNR')
    NF, MARGIN, reqSnr = 4, 8, 37.0
    model = {320: -40, 160: -42, 80: -44.5, 40: -47.5, 20: -50.5}
    rs.append(['SNR \u2192 RSSI \u2014 worked example (EHT, 0.8\u00b5s GI, MCS13)'])
    rs.append(['Noise Floor = \u2212174 + 10\u00b7log10(BW\u00d710^6) + Noise Figure + RSSI implementation '
               'Margin ; Estimated RSSI = Noise Floor + required SNR ; RSSI to SNR Diff = \u2212Noise Floor '
               '(add to RSSI to get SNR)'])
    rs.append([])
    rs.append(['Noise figure (dB)', 'RSSI implementation Margin (dB)', 'BW (MHz)', 'Noise Floor (dBm)',
               'required SNR (dB)', 'Estimated RSSI (dBm)', 'RSSI from model (dBm)', 'RSSI to SNR Diff (dB)'])
    for i in range(1, 9):
        c = rs.cell(4, i)
        c.fill = PatternFill('solid', fgColor='E2E8F0'); c.font = Font(bold=True)
        c.alignment = Alignment('center', 'center', wrap_text=True); c.border = BORD
    for i, b in enumerate([320, 160, 80, 40, 20]):
        row = 5 + i
        rs.append([NF if i == 0 else '', MARGIN if i == 0 else '', b,
                   f'=-174 + 10*LOG10(C{row}*10^6) + $A$5 + $B$5', reqSnr,
                   f'=D{row} + E{row}', model[b], f'=-D{row}'])
        for j in range(1, 9):
            rs.cell(row, j).alignment = Alignment(horizontal='center'); rs.cell(row, j).border = BORD
    for i, w in enumerate([16, 26, 10, 16, 15, 18, 18, 18], 1):
        rs.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w


def main():
    os.makedirs(OUT, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for bw in BANDWIDTHS:
        rows, off, src = run_bw(bw)
        write_tva_sheet(wb.create_sheet(f'TVA {bw}MHz'), bw, rows, off, src)
        write_charts(wb, bw, rows)
        matched = sum(1 for r in rows if r['lab'])
        print(f"{bw:>3} MHz -> TVA {bw}MHz / Charts {bw}MHz  "
              f"(offset {round(off,1)} dB, {matched}/{len(rows)} lab-matched)")
    write_rssi_sheet(wb)
    out = os.path.join(OUT, 'tva_sim_vs_lab_ALL.xlsx')
    wb.save(out)
    print('done ->', out)


if __name__ == '__main__':
    main()
