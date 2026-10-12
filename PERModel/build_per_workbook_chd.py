"""Human-readable workbook of ALL Channel-D (BF) PER permutations, from per_catalog_chd.npz.

Operating point = Format x FrameBW x ChannelBW x Antenna x Nss x GI x LTF x Control.
With 288 permutations, per-case chart sheets are impractical, so this emits two
compact summary sheets:

  * Sensitivity SNR@10%  - wide: the 8 dims + one column per MCS -> SNR(dB) at 10% PER
  * Index                - one row per measured curve (dims, MCS, sensitivity, sweep span)

Run from repo root (after build_per_catalog_chd.py):
    python PERModel/build_per_workbook_chd.py
"""
from __future__ import annotations

import os

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CAT = "PERModel/per_catalog_chd.npz"
OUT = "PERModel/PER_Model_ChD_AllPermutations.xlsx"
DIMS = ["Format", "Frame BW", "Channel BW", "Ant", "Nss", "GI (us)", "LTF", "Control"]


def _load():
    d = np.load(CAT)
    keys = [str(k) for k in d["keys"]]
    curve = {k: (d[f"x{i}"], d[f"y{i}"]) for i, k in enumerate(keys)}
    sens = {k: float(d["snr_sens"][i]) for i, k in enumerate(keys)}
    return keys, curve, sens


def build():
    keys, curve, sens = _load()
    parts = [k.split(",") for k in keys]
    mcs_all = sorted({int(p[8]) for p in parts})

    # unique 8-dim operating points, in first-seen order
    prefixes, seen = [], set()
    for p in parts:
        pre = tuple(p[:8])
        if pre not in seen:
            seen.add(pre)
            prefixes.append(pre)

    wb = Workbook()
    wb.remove(wb.active)
    hfill = PatternFill("solid", fgColor="1F2A44")
    hfont = Font(bold=True, color="FFFFFF")
    center = Alignment(horizontal="center")

    # Sensitivity SNR@10% (wide)
    ws = wb.create_sheet("Sensitivity SNR@10%")
    heads = DIMS + [f"MCS{m}" for m in mcs_all]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = hfont; c.fill = hfill; c.alignment = center
    for r, pre in enumerate(prefixes, start=2):
        for j, v in enumerate(pre, start=1):
            ws.cell(row=r, column=j, value=v)
        for j, m in enumerate(mcs_all, start=len(DIMS) + 1):
            k = ",".join(pre) + f",{m}"
            s = sens.get(k)
            if s is not None and not np.isnan(s):
                cell = ws.cell(row=r, column=j, value=round(s, 1))
                cell.number_format = "0.0"; cell.alignment = center
    ws.freeze_panes = "I2"
    for col, w in zip("ABCDEFGH", (7, 9, 11, 6, 5, 7, 5, 8)):
        ws.column_dimensions[col].width = w
    for j in range(len(DIMS) + 1, len(heads) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 7

    # Index (one row per curve)
    idx = wb.create_sheet("Index")
    iheads = DIMS + ["MCS", "Sensitivity SNR@10% (dB)", "Points", "SNR start", "SNR end"]
    for j, h in enumerate(iheads, start=1):
        c = idx.cell(row=1, column=j, value=h); c.font = hfont; c.fill = hfill
    for r, (p, k) in enumerate(zip(parts, keys), start=2):
        x, _ = curve[k]
        s = sens.get(k)
        vals = p[:8] + [int(p[8]),
                        None if s is None or np.isnan(s) else round(s, 1),
                        int(x.size), round(float(x.min()), 1), round(float(x.max()), 1)]
        for j, v in enumerate(vals, start=1):
            idx.cell(row=r, column=j, value=v)
    idx.freeze_panes = "A2"

    wb.save(OUT)
    print(f"Saved {OUT}: {len(prefixes)} permutations, {len(keys)} curves, "
          f"{os.path.getsize(OUT)/1024:.0f} KB, {len(wb.sheetnames)} sheets")


if __name__ == "__main__":
    build()
